"""Tests for the Markdown→docx renderer (subset, must never crash)."""

from __future__ import annotations

import re

from docx import Document

from app.summarize.docx_export import SummaryMeta, write_markdown_docx


def _meta(method: str | None = "Map-Reduce") -> SummaryMeta:
    return SummaryMeta(
        session_name="20260523_144550_Тест",
        date_str="23.05.2026 14:45",
        segment_count=7,
        speakers=["Я", "Собеседники"],
        duration_str="12:34",
        method=method,
    )


def _write(tmp_path, markdown: str, method: str | None = "Map-Reduce"):
    out = tmp_path / "out.docx"
    write_markdown_docx(markdown, _meta(method), out)
    return Document(str(out))


def _styles(doc):
    return [(p.style.name, p.text) for p in doc.paragraphs]


def test_headings_levels(tmp_path):
    doc = _write(tmp_path, "# Один\n## Два\n### Три\n#### Четыре\n")
    styles = dict((t, s) for s, t in _styles(doc))
    assert styles["Один"] == "Heading 1"
    assert styles["Два"] == "Heading 2"
    assert styles["Три"] == "Heading 3"
    # Deeper than ### degrades to Heading 3, never crashes.
    assert styles["Четыре"] == "Heading 3"


def test_bullets_and_indentation(tmp_path):
    md = "- один\n* два\n  - вложенный\n    - глубже\n"
    doc = _write(tmp_path, md)
    styles = dict((t, s) for s, t in _styles(doc))
    assert styles["один"] == "List Bullet"
    assert styles["два"] == "List Bullet"
    assert styles["вложенный"] == "List Bullet 2"
    assert styles["глубже"] == "List Bullet 2"


def test_numbered_list(tmp_path):
    doc = _write(tmp_path, "1. первое\n2) второе\n")
    styles = dict((t, s) for s, t in _styles(doc))
    assert styles["первое"] == "List Number"
    assert styles["второе"] == "List Number"


def test_bold_inline(tmp_path):
    doc = _write(tmp_path, "обычный **жирный** хвост\n")
    para = next(p for p in doc.paragraphs if "жирный" in p.text)
    runs = {r.text: r.bold for r in para.runs}
    assert runs["жирный"] is True
    assert runs["обычный "] is None
    assert runs[" хвост"] is None


def test_plain_paragraph_and_empty(tmp_path):
    doc = _write(tmp_path, "Просто абзац.\n")
    assert any(p.text == "Просто абзац." for p in doc.paragraphs)

    doc = _write(tmp_path, "   \n")
    assert any(p.text == "—" for p in doc.paragraphs)


def test_hrule_skipped(tmp_path):
    doc = _write(tmp_path, "до\n---\n***\nпосле\n")
    texts = [p.text for p in doc.paragraphs]
    assert "до" in texts and "после" in texts
    assert "---" not in texts and "***" not in texts


def test_code_fence_degrades_to_plain_paragraphs(tmp_path):
    doc = _write(tmp_path, "```\ncode line **не жирный**\n```\n")
    texts = [p.text for p in doc.paragraphs]
    assert "code line **не жирный**" in texts  # verbatim, no bold parsing
    assert "```" not in texts


def test_table_degrades_to_plain_text(tmp_path):
    doc = _write(tmp_path, "| a | b |\n| - | - |\n")
    texts = [p.text for p in doc.paragraphs]
    assert "| a | b |" in texts


def test_control_chars_are_stripped_never_crash(tmp_path):
    # C0 control characters are XML-illegal: without stripping, lxml raises
    # ValueError("All strings must be XML compatible"). U+FFFE/U+FFFF also
    # raise ValueError; a lone surrogate \ud800 raises UnicodeEncodeError.
    md = (
        "# Заголовок\x00\n"
        "- пункт\x01\n"
        "1. нумер\x02\n"
        "обычный **жирный\x03** хвост\n"
        "```\nкод\x04\n```\n"
        "хвосты \ufffe\uffff\ud800 конец\n"
    )
    doc = _write(tmp_path, md)  # must not raise
    texts = [p.text for p in doc.paragraphs]
    assert all(
        not re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", t)
        for t in texts
    )
    assert any("хвосты" in t and "конец" in t for t in texts)
    assert "Заголовок" in texts
    assert "пункт" in texts
    assert "нумер" in texts
    assert "код" in texts
    para = next(p for p in doc.paragraphs if "жирный" in p.text)
    runs = {r.text: r.bold for r in para.runs}
    assert runs["жирный"] is True


def test_control_chars_in_meta_are_stripped(tmp_path):
    meta = _meta()
    meta.session_name = "Сессия\x00"
    meta.date_str = "23.05.2026\x01 14:45"
    meta.method = "Map\x02-Reduce"
    out = tmp_path / "out.docx"
    write_markdown_docx("# Отчёт\n", meta, out)  # must not raise
    doc = Document(str(out))
    assert doc.paragraphs[0].text == "Протокол встречи — Сессия"
    assert "Дата: 23.05.2026 14:45" in doc.paragraphs[1].text
    assert "Метод: Map-Reduce" in doc.paragraphs[1].text


def test_weird_input_never_crashes(tmp_path):
    md = (
        "**\n"
        "###\n"
        "#### без пробела\n"
        "-\n"
        "* * * * *\n"
        "1.\n"
        "незакрытый **жирный\n"
        "«»„“‘’\n"
        "   -   пробелы   \n"
        "```незакрытый фенс\n"
        "строка после\n"
    )
    doc = _write(tmp_path, md)  # must not raise
    assert any("строка после" in p.text for p in doc.paragraphs)


def test_header_block(tmp_path):
    doc = _write(tmp_path, "# Отчёт\n")
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "Протокол встречи — 20260523_144550_Тест"
    meta = texts[1]
    assert "Дата: 23.05.2026 14:45" in meta
    assert "Реплик (после объединения): 7" in meta
    assert "Участники: Я, Собеседники" in meta
    assert "Длительность: 12:34" in meta
    assert "Метод: Map-Reduce" in meta


def test_header_without_method(tmp_path):
    doc = _write(tmp_path, "# Отчёт\n", method=None)
    meta = doc.paragraphs[1].text
    assert "Метод" not in meta


def test_empty_speakers_render_dash(tmp_path):
    meta = _meta()
    meta.speakers = []
    out = tmp_path / "out.docx"
    write_markdown_docx("# Отчёт\n", meta, out)
    doc = Document(str(out))
    assert "Участники: —" in doc.paragraphs[1].text
